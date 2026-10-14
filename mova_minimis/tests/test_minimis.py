"""Pruebas sin red: cálculo de saldo, validación de identificadores y permisos.

    python -m unittest discover -s tests -v
"""
import os
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
os.environ["MOVA_MINIMIS_DATOS"] = tempfile.mkdtemp(prefix="minimis_test_")

from logic import calculador, parser  # noqa: E402


def ayuda(fecha, importe, reglamento="REG (UE) 2023/2831 de minimis, General"):
    from logic.consultor import clase_reglamento
    return {"fecha_concesion": fecha, "importe_concedido": importe, "reglamento": reglamento,
            "clase_reglamento": clase_reglamento(reglamento), "codigo_concesion": "X"}


def res(*ayudas, completa=True, confianza="alta"):
    return {"cif": "B66894452", "razon_social": "PRUEBA SL", "ayudas_minimis_encontradas": list(ayudas),
            "busqueda_completa": completa, "confianza_busqueda": confianza}


class Calculo(unittest.TestCase):
    def setUp(self):
        self.per = calculador.periodo("movil", date(2026, 10, 13), None)

    def test_techo_y_ventana_movil(self):
        self.assertEqual(self.per["techo"], 300_000)
        c, ay = calculador.calcular_cliente(res(ayuda("2023-10-13", 50_000), ayuda("2023-10-14", 10_000)), self.per)
        # el 13/10/2023 ya está fuera (periodo = (13/10/2023, 13/10/2026]); el 14/10/2023 dentro
        self.assertEqual(c["total_concedido_3anos"], 10_000)
        self.assertEqual(c["saldo_disponible"], 290_000)
        self.assertEqual([a["computa"] for a in ay], [False, True])

    def test_estados(self):
        casos = [(100_000, "SALDO_DISPONIBLE"), (296_000, "LIMITE_ALCANZADO"), (300_000, "LIMITE_ALCANZADO"),
                 (300_001, "LIMITE_SUPERADO")]
        for importe, estado in casos:
            c, _ = calculador.calcular_cliente(res(ayuda("2025-01-01", importe)), self.per)
            self.assertEqual(c["estado"], estado, importe)

    def test_agricola_suma_al_techo_general_y_tiene_sublimite(self):
        c, _ = calculador.calcular_cliente(res(ayuda("2025-05-05", 200_000),
                                               ayuda("2025-06-06", 45_000, "REG (UE) 1408/2013 de minimis")), self.per)
        self.assertEqual(c["saldo_disponible"], 55_000)
        self.assertEqual(c["sublimites"][0]["disponible"], 5_000)

    def test_sieg_techo_combinado(self):
        c, _ = calculador.calcular_cliente(res(ayuda("2025-05-05", 100_000),
                                               ayuda("2025-06-06", 600_000, "REG (UE) 2023/2832 de minimis, SIEG")),
                                           self.per)
        self.assertEqual(c["saldo_disponible"], 50_000)  # 750.000 - 700.000 < 300.000 - 100.000

    def test_futura_no_computa(self):
        c, ay = calculador.calcular_cliente(res(ayuda("2026-12-01", 99_000)), self.per)
        self.assertEqual(c["total_concedido_3anos"], 0)
        self.assertFalse(ay[0]["computa"])

    def test_liberacion_y_recuperacion(self):
        c, _ = calculador.calcular_cliente(res(ayuda("2024-02-29", 250_000), ayuda("2025-01-10", 60_000)), self.per)
        self.assertEqual(c["estado"], "LIMITE_SUPERADO")
        self.assertEqual(c["fecha_recupera_saldo"], "2027-03-01")  # el 28/02/2027 aún computa

    def test_saldo_a_1_de_enero(self):
        # el 01/01/2027 el periodo es (01/01/2024, 01/01/2027]: salen las de 2023 y la del mismo 01/01/2024
        c, _ = calculador.calcular_cliente(res(ayuda("2023-12-01", 20_000), ayuda("2024-01-01", 50_000),
                                               ayuda("2024-01-02", 30_000)), self.per)
        self.assertEqual(c["saldo_disponible"], 200_000)
        self.assertEqual(c["fecha_saldo_enero"], "2027-01-01")
        self.assertEqual(c["saldo_1_enero"], 270_000)

    def test_saldo_a_1_de_enero_cuenta_las_ya_concedidas_despues(self):
        c, _ = calculador.calcular_cliente(res(ayuda("2026-12-01", 99_000)), self.per)
        self.assertEqual(c["saldo_disponible"], 300_000)
        self.assertEqual(c["saldo_1_enero"], 201_000)

    def test_saldo_a_1_de_enero_por_ejercicios(self):
        per = calculador.periodo("ejercicios", date(2026, 10, 13), None)  # 2023-2025 -> el 1 de enero, 2024-2026
        c, _ = calculador.calcular_cliente(res(ayuda("2023-05-05", 10_000), ayuda("2026-03-03", 5_000)), per)
        self.assertEqual(c["saldo_disponible"], 290_000)
        self.assertEqual(c["saldo_1_enero"], 295_000)

    def test_busqueda_incompleta_es_pendiente(self):
        c, _ = calculador.calcular_cliente(res(completa=False), self.per)
        self.assertEqual(c["estado"], "PENDIENTE")
        self.assertIsNone(c["saldo_disponible"])
        self.assertIsNone(c["saldo_1_enero"])

    def test_no_encontrado_tiene_saldo_pero_queda_pendiente(self):
        c, _ = calculador.calcular_cliente(res(confianza="no encontrado"), self.per)
        self.assertEqual(c["saldo_disponible"], 300_000)
        self.assertTrue(any("PENDIENTE" in p for p in c["pendientes"]))

    def test_ejercicios_antiguos_usan_1407(self):
        per = calculador.periodo("ejercicios", date(2026, 1, 1), [2021, 2022, 2023])
        self.assertEqual(per["techo"], 200_000)
        per = calculador.periodo("ejercicios", date(2026, 1, 1), [2025])
        self.assertTrue(any("incompleto" in a for a in per["avisos"]))


class Identificadores(unittest.TestCase):
    def test_validos(self):
        for v in ("A46103834", "B66894452", "ES-B66894452", "b 66.894.452", "12345678Z", "X1234567L", "P0800000B"):
            self.assertTrue(parser.validar_identificador(v)[2], v)

    def test_invalidos(self):
        for v in ("B12345678", "12345678A", "B1234", "ACME"):
            self.assertFalse(parser.validar_identificador(v)[2], v)

    def test_texto_libre(self):
        d, e = {}, []
        parser.extraer_de_texto("Cliente: CAL PETITO DE TORA SL, CIF B-66894452\nOtro B12345678", "t", d, e)
        self.assertIn("B66894452", d)
        self.assertEqual(d["B66894452"]["razon_social"], "CAL PETITO DE TORA SL")
        self.assertEqual(len(e), 1)

    def test_etiquetas_solo_con_dos_puntos(self):
        salida = parser.procesar_manual("K1234567L ; Cliente Uno SL\nRazón social: Talleres Ruiz SL - CIF: A46103834")
        nombres = {c["cif"]: c["razon_social"] for c in salida["clientes"]}
        self.assertEqual(nombres["K1234567L"], "Cliente Uno SL")
        self.assertEqual(nombres["A46103834"], "Talleres Ruiz SL")

    def test_tabla_con_cabecera(self):
        d, e = {}, []
        parser.extraer_de_tabla([["Razón social", "NIF"], ["MERCADONA SA", "A46103834"]], "x", d, e)
        self.assertEqual(d["A46103834"]["razon_social"], "MERCADONA SA")


class Regresiones(unittest.TestCase):
    """Fallos encontrados en la revisión del 13/10/2026."""

    @classmethod
    def setUpClass(cls):
        import app as aplicacion
        cls.app = aplicacion.crear_app()

    def test_numeros_de_norma_no_son_cif_ilegibles(self):
        d, e = {}, []
        parser.extraer_de_texto("Reglamento (UE) 2023/2831 y 1407/2013, expediente 2025/12345", "x", d, e)
        self.assertEqual(e, [])
        parser.extraer_de_texto("CIF: B636/9153", "x", d, e)
        self.assertEqual(len(e), 1)

    def test_redireccion_abierta(self):
        from routes.auth import _destino_seguro
        with self.app.test_request_context():
            for malo in ("/\\evil.com", "//evil.com", "https://evil.com", "/\\/evil.com"):
                self.assertEqual(_destino_seguro(malo), "/", malo)
            self.assertEqual(_destino_seguro("/consultas/nueva"), "/consultas/nueva")

    def test_auditar_usuario_inexistente(self):
        from models import db
        db.auditar(987654, "borrado@x", "restauracion")
        self.assertIsNone(db.listar_auditoria(1)[0]["usuario_id"])

    def test_poda_de_copias(self):
        import config
        from models import db
        for p in config.DIR_BACKUPS.glob("*.db"):
            p.unlink()
        for i in range(7):
            (config.DIR_BACKUPS / f"previa_restaurar_2026010{i}_000000.db").write_bytes(b"x")
        for i in range(12):
            (config.DIR_BACKUPS / f"minimis_202601{i:02d}_000000.db").write_bytes(b"x")
        db._podar(db.PREFIJO_PREVIA, 5)
        db.copia_automatica()
        self.assertEqual(len(list(config.DIR_BACKUPS.glob("minimis_*.db"))), 10)
        self.assertEqual(len(list(config.DIR_BACKUPS.glob("previa_restaurar_*.db"))), 5)

    def test_zip_restaurable_filtra_nombres(self):
        import io
        import zipfile
        from models import db
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("minimis.db", b"")
            z.writestr("documentos/" + "a" * 64 + ".pdf.enc", b"x")
            z.writestr("documentos/../../malo.exe", b"x")
            z.writestr("documentos/app.py", b"x")
        nombres = [i.filename for i in db.documentos_restaurables(zipfile.ZipFile(buf))]
        self.assertEqual(nombres, ["documentos/" + "a" * 64 + ".pdf.enc"])

    def test_ejercicios_libera_cada_reglamento(self):
        per = calculador.periodo("ejercicios", date(2026, 1, 1), [2023, 2024, 2025])
        c, _ = calculador.calcular_cliente(res(ayuda("2023-05-05", 100_000),
                                               ayuda("2023-06-06", 600_000, "REG (UE) 2023/2832 de minimis, SIEG"),
                                               ayuda("2025-01-01", 150_000)), per)
        self.assertEqual(c["saldo_disponible"], -100_000)  # 750.000 - 850.000
        self.assertEqual(c["liberaciones"][0]["saldo_tras_liberar"], 150_000)  # sale todo 2023


    def test_29_febrero_computa_hasta_el_28(self):
        per = calculador.periodo("movil", date(2027, 2, 28), None)
        _, ay = calculador.calcular_cliente(res(ayuda("2024-02-29", 1000)), per)
        self.assertTrue(ay[0]["computa"])
        per = calculador.periodo("movil", date(2027, 3, 1), None)
        _, ay = calculador.calcular_cliente(res(ayuda("2024-02-29", 1000)), per)
        self.assertFalse(ay[0]["computa"])

    def test_importes(self):
        from routes.resultados import _importe
        casos = {"150.000": 150000, "1.500": 1500, "1.234,56": 1234.56, "1234,5": 1234.5, "1234.56": 1234.56,
                 "12 500 €": 12500}
        for texto, valor in casos.items():
            self.assertEqual(_importe(texto), valor, texto)
        for malo in ("1.23.4", "1,234,5", "abc", "0", "-5"):
            with self.assertRaises(ValueError, msg=malo):
                _importe(malo)

    def test_reglamentos_modificados(self):
        from logic.consultor import clase_reglamento
        self.assertEqual(clase_reglamento("Reglamento (UE) 2024/3118"), "agricola")
        self.assertEqual(clase_reglamento("Reglamento (UE) 2023/2391"), "pesca")

    def test_nif_sin_ceros_y_klm(self):
        self.assertEqual(parser.validar_identificador("1234567L")[:3], ("01234567L", "NIF", True))
        self.assertTrue(parser.validar_identificador("K1234567L")[2])
        d, e = {}, []
        parser.extraer_de_texto("Fecha de concesión 1/02/2024", "x", d, e)
        self.assertEqual(e, [])


class Regresiones102(unittest.TestCase):
    """Segunda revisión (13/10/2026)."""

    @classmethod
    def setUpClass(cls):
        import app as aplicacion
        from auth import user
        from models import db
        cls.app = aplicacion.crear_app()
        cls.app.testing = True
        if not db.usuario_por_email("admin2@prueba"):
            db.crear_usuario("admin2@prueba", "Admin", user.hash_clave("ClaveAdmin123"), "admin", cambiar_clave=0)

    def _admin(self):
        c = self.app.test_client()
        c.get("/login")
        with c.session_transaction() as s:
            csrf = s["csrf"]
        c.post("/login", data={"email": "admin2@prueba", "clave": "ClaveAdmin123", "csrf": csrf})
        return c, csrf

    def test_zip_sin_base_no_da_500_y_no_deja_restos(self):
        import io
        import zipfile
        import config
        c, csrf = self._admin()
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("informe.pdf", b"%PDF")
        buf.seek(0)
        r = c.post("/admin/backups/restaurar", data={"csrf": csrf, "copia": (buf, "informes.zip")},
                   content_type="multipart/form-data")
        self.assertEqual(r.status_code, 302)
        self.assertEqual(list(config.DIR_TEMPORAL.iterdir()), [])

    def test_descarga_de_copia_no_deja_zip(self):
        import config
        c, _ = self._admin()
        r = c.get("/admin/backups/descargar")
        self.assertEqual(r.data[:2], b"PK")
        r.close()
        self.assertEqual(list(config.DIR_TEMPORAL.iterdir()), [])
        self.assertEqual(list(config.DIR_BACKUPS.glob("*.zip")), [])

    def test_klm_y_dni_corto_en_entrada_manual(self):
        salida = parser.procesar_manual("K1234567L ; Cliente uno\n1234567L Cliente dos")
        self.assertEqual(sorted(x["cif"] for x in salida["clientes"]), ["01234567L", "K1234567L"])
        self.assertEqual(salida["errores"], [])

    def test_sublimites_segun_fecha(self):
        self.assertEqual(calculador.sublimite("agricola", date(2024, 12, 16))[0], 50_000)
        self.assertEqual(calculador.sublimite("agricola", date(2024, 12, 15))[0], 25_000)
        self.assertEqual(calculador.sublimite("agricola", date(2024, 6, 1))[0], 20_000)
        self.assertEqual(calculador.sublimite("pesca", date(2024, 12, 2))[0], 40_000)
        self.assertEqual(calculador.sublimite("pesca", date(2024, 12, 1))[0], 30_000)
        per = calculador.periodo("movil", date(2024, 6, 1), None)
        c, _ = calculador.calcular_cliente(res(ayuda("2024-01-10", 30_000, "REG (UE) 1408/2013 de minimis")), per)
        self.assertEqual(c["sublimites"][0]["disponible"], -10_000)

    def test_copias_antiguas_no_rompen_la_poda(self):
        import config
        from models import db
        for p in config.DIR_BACKUPS.glob("*.db"):
            p.unlink()
        for i in range(12):
            (config.DIR_BACKUPS / f"minimis_antes_de_restaurar_202601{i:02d}_000000.db").write_bytes(b"x")
        nueva = db.copia_automatica()
        self.assertTrue(nueva.exists())

    def test_todas_las_pantallas_abren(self):
        """Una ruta borrada rompe cualquier plantilla que la enlace (url_for): se abre cada pantalla."""
        c, _ = self._admin()
        for ruta in ("/", "/consultas/nueva", "/admin/configuracion", "/admin/usuarios", "/admin/auditoria",
                     "/admin/backups", "/cambiar-clave"):
            self.assertEqual(c.get(ruta).status_code, 200, ruta)
        endpoints = {r.endpoint for r in self.app.url_map.iter_rules()}
        self.assertIn("admin.apagar", endpoints)

    def test_intentos_no_crean_entradas(self):
        from routes import auth
        auth._fallos.clear()
        self.assertFalse(auth._bloqueado(("127.0.0.1", "nadie@x")))
        self.assertEqual(auth._fallos, {})


class DesgloseMovil(unittest.TestCase):
    """Tres años hacia atrás desde la fecha del informe tocan cuatro años naturales; el primero y el último, en parte."""

    def test_tramos(self):
        per = calculador.periodo("movil", date(2026, 10, 13), None)
        self.assertEqual(per["etiqueta"], "14/10/2023 – 13/10/2026")
        c, _ = calculador.calcular_cliente(res(ayuda("2023-12-01", 1000), ayuda("2026-01-05", 2000)), per)
        tramos = {x["ano"]: x["tramo"] for x in c["desglose_por_ano"]}
        self.assertEqual(tramos, {2023: "desde el 14/10", 2024: "", 2025: "", 2026: "hasta el 13/10"})

    def test_fin_de_ano_sin_tramo(self):
        per = calculador.periodo("movil", date(2025, 12, 31), None)
        c, _ = calculador.calcular_cliente(res(), per)
        self.assertEqual([(x["ano"], x["tramo"]) for x in c["desglose_por_ano"]], [(2023, ""), (2024, ""), (2025, "")])


class Escritorio(unittest.TestCase):
    """Versión de escritorio: sin inicio de sesión, token obligatorio y candado en la administración."""

    TOKEN = "t" * 48

    @classmethod
    def setUpClass(cls):
        import app as aplicacion
        import config
        cls.config = config
        cls.antes = (config.ESCRITORIO, config.TOKEN_ESCRITORIO, config.PUERTO)
        config.ESCRITORIO, config.TOKEN_ESCRITORIO, config.PUERTO = True, cls.TOKEN, 5099
        cls.app = aplicacion.crear_app()
        cls.app.testing = True

    @classmethod
    def tearDownClass(cls):
        cls.config.ESCRITORIO, cls.config.TOKEN_ESCRITORIO, cls.config.PUERTO = cls.antes

    def _cliente(self, token=TOKEN, host="127.0.0.1:5099"):
        """Cliente que se presenta como la ventana: token en cabecera y la dirección que usa Electron."""
        c = self.app.test_client()
        c.environ_base["HTTP_X_MOVA_TOKEN"] = token
        original = c.open

        def abrir(*a, **kw):
            kw.setdefault("base_url", f"http://{host}")
            return original(*a, **kw)
        c.open = abrir
        return c

    def _csrf(self, c):
        c.get("/")
        with c.session_transaction(base_url="http://127.0.0.1:5099") as s:
            return s["csrf"]

    def test_sin_token_o_con_otro_host_no_entra(self):
        self.assertEqual(self._cliente(token="malo").get("/").status_code, 403)
        self.assertEqual(self._cliente(token="").get("/salud").status_code, 403)
        self.assertEqual(self._cliente(host="localhost:5099").get("/").status_code, 403)

    def test_entra_sin_contrasena(self):
        c = self._cliente()
        r = c.get("/")
        self.assertEqual(r.status_code, 200)
        self.assertIn(b"Consultas de minimis", r.data)
        self.assertEqual(c.get("/consultas/nueva").status_code, 200)
        self.assertEqual(c.get("/login").status_code, 302)

    def test_administracion_con_candado(self):
        from auth import user
        from models import db
        from routes.auth import CLAVE_CONFIG
        # Sin ajuste en la base manda la contraseña de serie (aquí, una de prueba en su lugar)
        db.guardar_ajuste(CLAVE_CONFIG, None)
        antes = self.config.CLAVE_CONFIG_DEFECTO
        self.config.CLAVE_CONFIG_DEFECTO = user.hash_clave("ConfigSegura123")
        self.addCleanup(setattr, self.config, "CLAVE_CONFIG_DEFECTO", antes)
        c = self._cliente()
        csrf = self._csrf(c)
        for ruta in ("/admin/configuracion", "/admin/backups", "/admin/auditoria"):
            r = c.get(ruta)
            self.assertEqual(r.status_code, 302, ruta)
            self.assertIn("/desbloquear", r.headers["Location"])
        self.assertNotIn(b"repetir", c.get("/desbloquear").data)  # no hay que crearla: viene puesta
        r = c.post("/desbloquear?next=/admin/backups", data={"csrf": csrf, "clave": "ConfigSegura123"})
        self.assertEqual(r.headers["Location"], "/admin/backups")
        self.assertEqual(c.get("/admin/backups").status_code, 200)
        self.assertEqual(c.get("/admin/usuarios").status_code, 302)  # sin gestión de usuarios
        # Bloquear y volver a abrir: la contraseña mala no vale, la buena sí
        c.post("/bloquear", data={"csrf": csrf})
        self.assertEqual(c.get("/admin/configuracion").status_code, 302)
        c.post("/desbloquear", data={"csrf": csrf, "clave": "Mala1234567"})
        self.assertEqual(c.get("/admin/configuracion").status_code, 302)
        c.post("/desbloquear", data={"csrf": csrf, "clave": "ConfigSegura123"})
        self.assertEqual(c.get("/admin/configuracion").status_code, 200)

    def test_cambiarla_sustituye_a_la_de_serie(self):
        from auth import user
        from models import db
        from routes.auth import CLAVE_CONFIG
        db.guardar_ajuste(CLAVE_CONFIG, None)
        antes = self.config.CLAVE_CONFIG_DEFECTO
        self.config.CLAVE_CONFIG_DEFECTO = user.hash_clave("DeSerie12345")
        self.addCleanup(setattr, self.config, "CLAVE_CONFIG_DEFECTO", antes)
        c = self._cliente()
        csrf = self._csrf(c)
        c.post("/desbloquear", data={"csrf": csrf, "clave": "DeSerie12345"})
        c.post("/admin/configuracion", data={"csrf": csrf, "accion": "cambiar_clave_config", "actual": "DeSerie12345",
                                             "nueva": "Propia123456", "repetir": "Propia123456"})
        c.post("/bloquear", data={"csrf": csrf})
        c.post("/desbloquear", data={"csrf": csrf, "clave": "DeSerie12345"})
        self.assertEqual(c.get("/admin/configuracion").status_code, 302)  # la de serie ya no vale
        c.post("/desbloquear", data={"csrf": csrf, "clave": "Propia123456"})
        self.assertEqual(c.get("/admin/configuracion").status_code, 200)

    def test_la_de_serie_no_esta_en_claro(self):
        import config
        self.assertTrue(config.CLAVE_CONFIG_DEFECTO.startswith("pbkdf2_sha256$"))

    def test_otra_sesion_no_hereda_el_desbloqueo(self):
        from models import db
        from auth import user
        from routes.auth import CLAVE_CONFIG
        db.guardar_ajuste(CLAVE_CONFIG, user.hash_clave("ConfigSegura123"))
        a = self._cliente()
        csrf = self._csrf(a)
        a.post("/desbloquear", data={"csrf": csrf, "clave": "ConfigSegura123"})
        self.assertEqual(a.get("/admin/configuracion").status_code, 200)
        self.assertEqual(self._cliente().get("/admin/configuracion").status_code, 302)


class Documentos(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())

    def test_excel_word_y_pdf(self):
        from docx import Document
        from openpyxl import Workbook
        import pymupdf
        wb = Workbook()
        wb.active.append(["Cliente", "CIF/NIF", "Año"])
        wb.active.append(["MERCADONA SA", "A46103834", 2025])
        wb.save(self.dir / "c.xlsx")
        d = Document()
        d.add_paragraph("La empresa CAL PETITO DE TORA SL con CIF B66894452 solicita la ayuda.")
        d.save(self.dir / "c.docx")
        pdf = pymupdf.open()
        pdf.new_page().insert_text((72, 72), "Beneficiario: MOSQUITO BOOKS BARCELONA SL - NIF: B17884859")
        pdf.save(self.dir / "c.pdf")
        for nombre, cif in (("c.xlsx", "A46103834"), ("c.docx", "B66894452"), ("c.pdf", "B17884859")):
            salida = parser.procesar_documento(self.dir / nombre, nombre)
            self.assertEqual([c["cif"] for c in salida["clientes"]], [cif], nombre)

    def test_imagen_sin_ocr_ni_claude_avisa(self):
        from PIL import Image
        Image.new("RGB", (50, 50), "white").save(self.dir / "s.png")
        if parser.tesseract_disponible():
            self.skipTest("Tesseract instalado")
        salida = parser.procesar_documento(self.dir / "s.png", "s.png")
        self.assertIn("OCR no disponible", salida["avisos"][0])


class Orquestacion(unittest.TestCase):
    """La cadena entera en el mismo hilo, con la BDNS sustituida por datos fijos."""

    @classmethod
    def setUpClass(cls):
        import app as aplicacion
        from integrations import web_scraper
        from models import db
        aplicacion.crear_app()
        cls.db = db
        cls.ws = web_scraper
        cls.originales = (web_scraper.buscar_beneficiarios, web_scraper.concesiones_minimis,
                          web_scraper.ficha_convocatoria)
        web_scraper.buscar_beneficiarios = lambda t, c=None: [{"id": 1, "cif": t, "nombre": "PRUEBA SL"}]
        web_scraper.concesiones_minimis = lambda i, c=None: [{
            "codigoConcesion": "SB1", "fechaConcesion": "2025-05-05", "fechaRegistro": "2025-06-01",
            "convocante": "ESTADO", "reglamento": "REG (UE) 2023/2831 de minimis, General",
            "instrumento": "SUBVENCIÓN", "numeroConvocatoria": "1", "ayudaEquivalente": 120000}]
        web_scraper.ficha_convocatoria = lambda n, c=None: {"titulo": "Programa de prueba", "url_bases": None}
        cls.uid = db.usuario_por_email("admin@local")["id"]

    @classmethod
    def tearDownClass(cls):
        (cls.ws.buscar_beneficiarios, cls.ws.concesiones_minimis, cls.ws.ficha_convocatoria) = cls.originales

    def _consulta(self, cif):
        import uuid
        cid = str(uuid.uuid4())
        self.db.crear_consulta(cid, self.uid, "movil", "2026-10-13", [], {
            "clientes": [{"cif": cif, "razon_social": "Prueba SL"}], "errores": [], "avisos": []})
        return cid

    def test_completa_y_exporta(self):
        from routes import orquestador
        cid = self._consulta("B66894452")
        orquestador._ejecutar(cid)
        q = self.db.consulta(cid)
        self.assertEqual(q["estado"], "completada")
        r = self.db.resultados_de(cid)[0]
        self.assertEqual(r["saldo_disponible"], 180_000)
        self.assertEqual(r["ayudas"][0]["programa"], "Programa de prueba")
        # Sin abrir Word (tarda un minuto la primera vez): el .docx sale del motor de la plantilla MOVA
        # y el PDF se prueba con el ejecutable. Aquí: Word, Excel y JSON, y el contenido del informe.
        from docx import Document
        from integrations import word_pdf
        antes = word_pdf._disponible
        word_pdf._disponible = False
        self.addCleanup(setattr, word_pdf, "_disponible", antes)
        for f in ("docx", "xlsx", "json"):
            self.assertTrue(orquestador.exportar(cid, f, {"id": self.uid}).exists(), f)
        ruta = orquestador.exportar(cid, "docx", {"id": self.uid})
        texto = " ".join(p.text for p in Document(ruta).paragraphs)
        self.assertIn("RESUMEN EJECUTIVO", texto)        # título de nivel 1 de la plantilla MOVA
        self.assertNotIn("Referencias de fuente", texto)
        # quitado a petición del usuario (1.1.4); el [PENDIENTE: única empresa] se queda
        self.assertNotIn("infosubvenciones.es", texto)
        self.assertNotIn("La BDNS registra", texto)
        self.assertIn("única empresa", texto)
        celdas = [c.text for t in Document(ruta).tables for f in t.rows for c in f.cells]
        self.assertNotIn(cid, celdas)  # sin la fila «Consulta» de la portada
        self.assertTrue(any(c.startswith("Saldo a 01/01/") for c in celdas))
        import openpyxl
        libro = openpyxl.load_workbook(orquestador.exportar(cid, "xlsx", {"id": self.uid}))
        valores = [c.value for f in libro["Resumen"].iter_rows() for c in f]
        self.assertNotIn(cid, valores)
        self.assertNotIn("Consulta", valores)
        # una consulta de antes de la 1.1.4 se pone al día al abrirla, sin volver a la BDNS
        import json
        viejo = json.loads(self.db.consulta(cid)["informe_json"])
        del viejo["informe"]["etiqueta_saldo_enero"]
        self.db.actualizar_consulta(cid, informe_json=json.dumps(viejo))
        self.assertIn("etiqueta_saldo_enero", orquestador.informe_al_dia(cid)["informe"])
        self.assertIn("TOTAL", celdas)
        self.assertIn("General · 2023/2831", celdas)

    def test_informe_sin_ayudas_que_no_computan(self):
        from logic import exportador
        inf = {"informe": {
            "titulo": "Informe", "reglamento": "Reglamento (UE) 2023/2831", "umbral": 300000,
            "periodo_consulta": "14/10/2023 – 13/10/2026", "fecha_generacion": "2026-10-13T10:00:00Z",
            "usuario": "Prueba", "audit_trail": {"consulta_id": "x"}, "errores_entrada": [],
            "notas_importantes": ["[PENDIENTE: única empresa] nota"],
            "resumen_ejecutivo": {"clientes_consultados": 1, "clientes_con_saldo": 1, "clientes_en_limite": 0,
                                  "clientes_superado_limite": 0, "clientes_pendientes": 0,
                                  "total_saldo_disponible_conjunto": 290000, "total_concedido_periodo": 10000,
                                  "total_limite_superado": 0},
            "detalle_clientes": [{
                "razon_social": "PRUEBA | SL", "cif": "B66894452", "saldo_disponible": 290000, "estado_texto": "Saldo disponible",
                "total_concedido": 10000, "desglose_por_ano": [], "proxima_actualizacion_limite": None,
                "fecha_recupera_saldo": None, "sublimites": [], "pendientes": [], "notas": "", "notas_busqueda": "",
                "referencias_fuente": [
                    {"fecha_concesion": "01/01/2025", "programa": "Computa", "convocante": "ESTADO", "computa": True,
                     "reglamento": "REG (UE) 2023/2831 de minimis, General", "clase": "General", "importe": 10000},
                    {"fecha_concesion": "01/01/2021", "programa": "Antigua", "convocante": "ESTADO", "computa": False,
                     "reglamento": "REG (UE) 1407/2013 de minimis", "clase": "General", "importe": 5000}]}]}}
        inf["informe"].update({"etiqueta_saldo": "Saldo a 13/10/2026", "etiqueta_saldo_enero": "Saldo a 01/01/2027"})
        inf["informe"]["resumen_ejecutivo"]["total_saldo_1_enero_conjunto"] = 295000
        inf["informe"]["detalle_clientes"][0]["saldo_1_enero"] = 295000
        md = exportador.markdown_informe(inf)
        self.assertIn("Saldo a 01/01/2027", md)
        self.assertIn("295.000,00 €", md)
        self.assertNotIn("consulta:", md)  # el identificador de la consulta no sale en la portada
        self.assertIn("Computa · ESTADO", md)
        self.assertNotIn("Antigua", md)
        self.assertIn("| TOTAL |", md)
        self.assertIn("PRUEBA / SL", md)  # la barra vertical rompería la tabla Markdown

    def test_cancelada(self):
        from routes import orquestador
        cid = self._consulta("A46103834")
        self.db.pedir_cancelacion(cid)
        orquestador._ejecutar(cid)
        self.assertEqual(self.db.consulta(cid)["estado"], "cancelada")

    def test_interrumpida_al_arrancar(self):
        cid = self._consulta("A46103834")
        self.db.marcar_interrumpidas()
        self.assertEqual(self.db.consulta(cid)["estado"], "interrumpida")


class Permisos(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import app as aplicacion
        from auth import user
        from models import db
        cls.app = aplicacion.crear_app()
        cls.app.testing = True
        db.crear_usuario("lector@prueba", "Lector", user.hash_clave("ClaveLector123"), "viewer", cambiar_clave=0)

    def _entrar(self, email, clave):
        c = self.app.test_client()
        c.get("/login")
        with c.session_transaction() as s:
            csrf = s["csrf"]
        r = c.post("/login", data={"email": email, "clave": clave, "csrf": csrf})
        return c, csrf, r

    def test_sin_sesion_redirige(self):
        r = self.app.test_client().get("/consultas/nueva")
        self.assertEqual(r.status_code, 302)
        self.assertIn("/login", r.headers["Location"])

    def test_admin_inicial_obligado_a_cambiar(self):
        c, _, r = self._entrar("admin@local", "TemporalAdmin123")
        self.assertIn("/cambiar-clave", r.headers["Location"])
        self.assertIn("/cambiar-clave", c.get("/").headers["Location"])

    def test_lector_no_consulta_ni_administra(self):
        c, csrf, _ = self._entrar("lector@prueba", "ClaveLector123")
        self.assertEqual(c.get("/").status_code, 200)
        self.assertEqual(c.get("/consultas/nueva").status_code, 403)
        self.assertEqual(c.get("/admin/usuarios").status_code, 403)
        r = c.post("/api/consultas", json={"clientes": [{"cif": "A46103834"}]}, headers={"X-CSRF": csrf})
        self.assertEqual(r.status_code, 403)

    def test_desactivado_queda_fuera_al_instante(self):
        from models import db
        c, _, _ = self._entrar("lector@prueba", "ClaveLector123")
        u = db.usuario_por_email("lector@prueba")
        db.actualizar_usuario(u["id"], activo=0)
        try:
            self.assertEqual(c.get("/").status_code, 302)
        finally:
            db.actualizar_usuario(u["id"], activo=1)

    def test_post_sin_csrf_rechazado(self):
        c, _, _ = self._entrar("lector@prueba", "ClaveLector123")
        self.assertEqual(c.post("/logout").status_code, 400)


if __name__ == "__main__":
    unittest.main()
