"""Agente 3b · Exportador a Word, PDF y Excel.

Determinista (python-docx, reportlab, openpyxl): no necesita IA y así el
documento es idéntico cada vez. El documento va en claro, porque se imprime o se
envía; la marca va en los títulos dorados y en las cabeceras de tabla oscuras.
Sin macros ni contraseñas.
"""
from datetime import datetime
from pathlib import Path

import config
from logic.calculador import euros

DORADO = "B8892F"       # --gold-d: legible sobre blanco (el #d9b45b se pierde en papel)
DORADO_CLARO = "D9B45B"
OSCURO = "141518"       # --card
BORDE = "2A2D33"
VERDE, AMBAR, ROJO, GRIS = "16A34A", "D97706", "DC2626", "6B7280"
COLOR_ESTADO = {"SALDO_DISPONIBLE": VERDE, "LIMITE_ALCANZADO": AMBAR, "LIMITE_SUPERADO": ROJO, "PENDIENTE": GRIS}
LEYENDA = "Datos extraídos de fuentes públicas. Requiere validación oficial."
LOGO = config.RAIZ_CODIGO / "static" / "marca" / "mova-ia-negro.png"


def _fecha_local(iso: str) -> str:
    try:
        return datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone().strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return iso


def nombre_base(informe: dict) -> str:
    marca = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"informe_minimis_{marca}_{informe['audit_trail']['consulta_id'][:8]}"


# Word, PDF y Excel solo enseñan las ayudas que computan en el periodo, sin la lista de referencias de fuente
_SIN_AYUDAS = "Sin ayudas de minimis que computen en el periodo."


def _computan(d: dict) -> list[dict]:
    return [x for x in d["referencias_fuente"] if x["computa"]]


def _kpis(inf: dict) -> list[tuple[str, str]]:
    r = inf["resumen_ejecutivo"]
    return [
        ("Clientes consultados", str(r["clientes_consultados"])),
        ("Con saldo disponible", str(r["clientes_con_saldo"])),
        ("En el límite (≤ 5.000 €)", str(r["clientes_en_limite"])),
        ("Límite superado", str(r["clientes_superado_limite"])),
        ("Con datos pendientes", str(r["clientes_pendientes"])),
        ("Saldo disponible conjunto", euros(r["total_saldo_disponible_conjunto"])),
        ("Total concedido en el periodo", euros(r["total_concedido_periodo"])),
        ("Exceso sobre el techo", euros(r["total_limite_superado"])),
    ]


# ======================================================================= WORD
def exportar_docx(informe: dict, destino: Path) -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    inf = informe["informe"]
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(1.8)

    est = doc.styles["Normal"]
    est.font.name = "Calibri"
    est.font.size = Pt(10)
    for nombre, tam in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        h = doc.styles[nombre]
        h.font.name = "Calibri"
        h.font.size = Pt(tam)
        h.font.bold = True
        h.font.color.rgb = RGBColor.from_string(DORADO)

    def sombrear(celda, color):
        tc = celda._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        tc.append(shd)

    def tabla(cabeceras, filas, anchos=None, tam=8.5):
        t = doc.add_table(rows=1, cols=len(cabeceras))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, c in enumerate(cabeceras):
            cel = t.rows[0].cells[i]
            cel.text = ""
            run = cel.paragraphs[0].add_run(c)
            run.bold = True
            run.font.size = Pt(tam)
            run.font.color.rgb = RGBColor.from_string(DORADO_CLARO)
            sombrear(cel, OSCURO)
        for fila in filas:
            celdas = t.add_row().cells
            for i, v in enumerate(fila):
                color = None
                if isinstance(v, tuple):
                    v, color = v
                celdas[i].text = ""
                run = celdas[i].paragraphs[0].add_run("" if v is None else str(v))
                run.font.size = Pt(tam)
                if color:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(color)
        if anchos:
            for fila in t.rows:
                for i, w in enumerate(anchos):
                    fila.cells[i].width = Cm(w)
        return t

    # Pie con leyenda, usuario y número de página
    pie = sec.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = pie.add_run(f"{LEYENDA} · Generado el {_fecha_local(inf['fecha_generacion'])} por {inf['usuario']} · Página ")
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor.from_string(GRIS)
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        run = pie.add_run()
        run.font.size = Pt(7.5)
        if tipo:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), tipo)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = texto
            run._r.append(instr)

    # Portada
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run().add_picture(str(LOGO), width=Cm(6))
    for _ in range(5):
        doc.add_paragraph()
    t = doc.add_paragraph()
    run = t.add_run(inf["titulo"])
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(DORADO)
    doc.add_paragraph()
    for etiqueta, valor in (("Periodo de consulta", inf["periodo_consulta"]),
                            ("Normativa", f"{inf['reglamento']} · techo {euros(inf['umbral'])}"),
                            ("Consultado por", inf["usuario"]),
                            ("Fecha de generación", _fecha_local(inf["fecha_generacion"])),
                            ("Identificador de consulta", inf["audit_trail"]["consulta_id"])):
        p = doc.add_paragraph()
        a = p.add_run(f"{etiqueta}: ")
        a.bold = True
        p.add_run(str(valor))
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(LEYENDA)
    r.italic = True
    r.font.color.rgb = RGBColor.from_string(GRIS)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Resumen
    doc.add_heading("Resumen ejecutivo", level=1)
    tabla(["Indicador", "Valor"], _kpis(inf), anchos=[9, 6], tam=9.5)

    doc.add_heading("Detalle por cliente", level=1)
    tabla(["CIF", "Razón social", "Saldo disponible", "Estado", "Años", "Total computado"],
          [(d["cif"], d["razon_social"],
            euros(d["saldo_disponible"]) if d["saldo_disponible"] is not None else "[PENDIENTE]",
            (d["estado_texto"], COLOR_ESTADO[d["estado"]]),
            ", ".join(map(str, d["anos_consultados"])),
            euros(d["total_concedido"]) if d["total_concedido"] is not None else "[PENDIENTE]")
           for d in inf["detalle_clientes"]],
          anchos=[2.3, 5.2, 2.6, 2.4, 2.2, 2.6])

    for d in inf["detalle_clientes"]:
        doc.add_heading(f"{d['razon_social']} · {d['cif']}", level=2)
        p = doc.add_paragraph()
        p.add_run("Estado: ").bold = True
        r = p.add_run(d["estado_texto"])
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(COLOR_ESTADO[d["estado"]])
        p.add_run(f"   ·   Saldo disponible: {euros(d['saldo_disponible'])}"
                  f"   ·   Computado en el periodo: {euros(d['total_concedido'])}")
        if d["desglose_por_ano"]:
            p = doc.add_paragraph()
            p.add_run("Desglose por año: ").bold = True
            p.add_run("; ".join(f"{x['ano']}{' (' + x['tramo'] + ')' if x.get('tramo') else ''}: {euros(x['importe_concedido'])}"
            for x in d["desglose_por_ano"]))
        if d["proxima_actualizacion_limite"]:
            p = doc.add_paragraph()
            p.add_run("Próxima salida de una ayuda del periodo: ").bold = True
            p.add_run(d["proxima_actualizacion_limite"])
        if d["fecha_recupera_saldo"]:
            p = doc.add_paragraph()
            p.add_run("Vuelve a tener saldo a partir del: ").bold = True
            p.add_run(d["fecha_recupera_saldo"])
        for s in d["sublimites"]:
            doc.add_paragraph(f"Sublímite {s['clase']} ({s.get('norma', '')}): computado {euros(s['computado'])} de {euros(s['tope'])}; "
                              f"disponible {euros(s['disponible'])}.")
        computan = _computan(d)
        if computan:
            tabla(["Fecha", "Convocatoria / programa", "Concedente", "Reglamento", "Importe"],
                  [(x["fecha_concesion"], x["programa"], x["convocante"], x["reglamento"], euros(x["importe"]))
                   for x in computan],
                  anchos=[1.9, 6.4, 3.6, 2.9, 2.4], tam=7.5)
        else:
            doc.add_paragraph(_SIN_AYUDAS)
        for pend in d["pendientes"]:
            p = doc.add_paragraph()
            r = p.add_run(pend)
            r.bold = True
            r.font.color.rgb = RGBColor.from_string(AMBAR)
        if d["notas"]:
            doc.add_paragraph(d["notas"])
        if d["notas_busqueda"]:
            p = doc.add_paragraph()
            r = p.add_run("Búsqueda: " + d["notas_busqueda"])
            r.italic = True
            r.font.size = Pt(8.5)


    if inf.get("errores_entrada"):
        doc.add_heading("Entradas descartadas", level=1)
        tabla(["Entrada", "Motivo"], [(e["entrada"], e["error"]) for e in inf["errores_entrada"]], anchos=[8, 9])

    doc.add_heading("Notas importantes", level=1)
    for n in inf["notas_importantes"]:
        doc.add_paragraph(n, style="List Bullet")

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destino)
    return destino


# ======================================================================= PDF
def exportar_pdf(informe: dict, destino: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)
    from xml.sax.saxutils import escape

    inf = informe["informe"]
    hex_ = lambda h: colors.HexColor("#" + h)  # noqa: E731
    base = getSampleStyleSheet()
    st = {
        "n": ParagraphStyle("n", parent=base["Normal"], fontName="Helvetica", fontSize=9, leading=12),
        "p": ParagraphStyle("p", parent=base["Normal"], fontName="Helvetica", fontSize=7.5, leading=9.5),
        "pc": ParagraphStyle("pc", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=7.5, leading=9.5,
                             textColor=hex_(DORADO_CLARO)),
        "h1": ParagraphStyle("h1", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=15,
                             textColor=hex_(DORADO), spaceBefore=10, spaceAfter=6),
        "h2": ParagraphStyle("h2", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=11.5,
                             textColor=hex_(DORADO), spaceBefore=10, spaceAfter=4),
        "tit": ParagraphStyle("tit", parent=base["Title"], fontName="Helvetica-Bold", fontSize=24, leading=29,
                              textColor=hex_(DORADO), alignment=0),
        "ley": ParagraphStyle("ley", parent=base["Normal"], fontName="Helvetica-Oblique", fontSize=9,
                              textColor=hex_(GRIS)),
        "pend": ParagraphStyle("pend", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=8.5,
                               textColor=hex_(AMBAR)),
    }
    P = lambda t, s="n": Paragraph(escape("" if t is None else str(t)), st[s])  # noqa: E731

    def tabla(cab, filas, anchos):
        datos = [[Paragraph(escape(c), st["pc"]) for c in cab]]
        estilos = [("BACKGROUND", (0, 0), (-1, 0), hex_(OSCURO)),
                   ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c9c9c9")),
                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                   ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f6f3ea")])]
        for i, f in enumerate(filas, start=1):
            fila = []
            for j, v in enumerate(f):
                if isinstance(v, tuple):
                    v, color = v
                    fila.append(Paragraph(f'<font color="#{color}"><b>{escape(str(v))}</b></font>', st["p"]))
                else:
                    fila.append(P(v, "p"))
            datos.append(fila)
        t = Table(datos, colWidths=[a * cm for a in anchos], repeatRows=1)
        t.setStyle(TableStyle(estilos))
        return t

    def pie(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(hex_(GRIS))
        canvas.drawCentredString(A4[0] / 2, 1.0 * cm,
                                 f"{LEYENDA} · {_fecha_local(inf['fecha_generacion'])} · {inf['usuario']} · "
                                 f"Página {doc_.page}")
        canvas.setStrokeColor(hex_(DORADO))
        canvas.setLineWidth(0.6)
        canvas.line(2 * cm, A4[1] - 1.3 * cm, A4[0] - 2 * cm, A4[1] - 1.3 * cm)
        canvas.restoreState()

    h = []
    if LOGO.exists():
        h.append(Image(str(LOGO), width=6 * cm, height=6 * cm * 137 / 640, hAlign="LEFT"))
    h += [Spacer(1, 5 * cm), Paragraph(escape(inf["titulo"]), st["tit"]), Spacer(1, 0.8 * cm)]
    for etq, val in (("Periodo de consulta", inf["periodo_consulta"]),
                     ("Normativa", f"{inf['reglamento']} · techo {euros(inf['umbral'])}"),
                     ("Consultado por", inf["usuario"]),
                     ("Fecha de generación", _fecha_local(inf["fecha_generacion"])),
                     ("Identificador de consulta", inf["audit_trail"]["consulta_id"])):
        h.append(Paragraph(f"<b>{escape(etq)}:</b> {escape(str(val))}", st["n"]))
    h += [Spacer(1, 4 * cm), Paragraph(escape(LEYENDA), st["ley"]), PageBreak()]

    h.append(Paragraph("Resumen ejecutivo", st["h1"]))
    h.append(tabla(["Indicador", "Valor"], _kpis(inf), [9, 6]))
    h.append(Paragraph("Detalle por cliente", st["h1"]))
    h.append(tabla(["CIF", "Razón social", "Saldo disponible", "Estado", "Años", "Total computado"],
                   [(d["cif"], d["razon_social"], euros(d["saldo_disponible"]),
                     (d["estado_texto"], COLOR_ESTADO[d["estado"]]),
                     ", ".join(map(str, d["anos_consultados"])), euros(d["total_concedido"]))
                    for d in inf["detalle_clientes"]],
                   [2.2, 5.0, 2.6, 2.5, 2.2, 2.5]))

    for d in inf["detalle_clientes"]:
        bloque = [Paragraph(escape(f"{d['razon_social']} · {d['cif']}"), st["h2"]),
                  Paragraph(f'<b>Estado:</b> <font color="#{COLOR_ESTADO[d["estado"]]}"><b>'
                            f'{escape(d["estado_texto"])}</b></font> · <b>Saldo disponible:</b> '
                            f'{escape(euros(d["saldo_disponible"]))} · <b>Computado:</b> '
                            f'{escape(euros(d["total_concedido"]))}', st["n"])]
        if d["desglose_por_ano"]:
            bloque.append(P("Desglose por año: " + "; ".join(
                f"{x['ano']}{' (' + x['tramo'] + ')' if x.get('tramo') else ''}: {euros(x['importe_concedido'])}"
            for x in d["desglose_por_ano"])))
        if d["proxima_actualizacion_limite"]:
            bloque.append(P(f"Próxima salida de una ayuda del periodo: {d['proxima_actualizacion_limite']}"))
        if d["fecha_recupera_saldo"]:
            bloque.append(P(f"Vuelve a tener saldo a partir del: {d['fecha_recupera_saldo']}"))
        h.append(KeepTogether(bloque))
        for s in d["sublimites"]:
            h.append(P(f"Sublímite {s['clase']} ({s.get('norma', '')}): computado {euros(s['computado'])} de {euros(s['tope'])}; "
                       f"disponible {euros(s['disponible'])}."))
        computan = _computan(d)
        if computan:
            h.append(Spacer(1, 4))
            h.append(tabla(["Fecha", "Convocatoria / programa", "Concedente", "Reglamento", "Importe"],
                           [(x["fecha_concesion"], x["programa"], x["convocante"], x["reglamento"], euros(x["importe"]))
                            for x in computan],
                           [1.8, 6.3, 3.6, 2.9, 2.4]))
        else:
            h.append(P(_SIN_AYUDAS))
        for pend in d["pendientes"]:
            h.append(P(pend, "pend"))
        if d["notas"]:
            h.append(P(d["notas"]))
        if d["notas_busqueda"]:
            h.append(P("Búsqueda: " + d["notas_busqueda"], "p"))

    if inf.get("errores_entrada"):
        h.append(Paragraph("Entradas descartadas", st["h1"]))
        h.append(tabla(["Entrada", "Motivo"], [(e["entrada"], e["error"]) for e in inf["errores_entrada"]], [8, 9]))
    h.append(Paragraph("Notas importantes", st["h1"]))
    for n in inf["notas_importantes"]:
        h.append(Paragraph("• " + escape(n), st["n"]))
        h.append(Spacer(1, 3))

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(destino), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm,
                            bottomMargin=1.8 * cm, title=inf["titulo"], author="MOVA IA · MOVA Minimis")
    doc.build(h, onFirstPage=pie, onLaterPages=pie)
    return destino


# ======================================================================= EXCEL
def exportar_xlsx(informe: dict, destino: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    inf = informe["informe"]
    wb = Workbook()
    moneda = '#,##0.00 "€";[Red]-#,##0.00 "€"'
    cab_fill = PatternFill("solid", fgColor=OSCURO)
    cab_font = Font(bold=True, color=DORADO_CLARO)
    fino = Side(style="thin", color="C9C9C9")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)

    def cabecera(ws, fila, textos, anchos):
        for i, (t, w) in enumerate(zip(textos, anchos), start=1):
            c = ws.cell(row=fila, column=i, value=t)
            c.fill, c.font, c.border = cab_fill, cab_font, borde
            c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = ws.cell(row=fila + 1, column=1)

    # Hoja 1 · Resumen
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = inf["titulo"]
    ws["A1"].font = Font(bold=True, size=14, color=DORADO)
    meta = [("Periodo de consulta", inf["periodo_consulta"]), ("Normativa", inf["reglamento"]),
            ("Techo aplicado", inf["umbral"]), ("Consultado por", inf["usuario"]),
            ("Fecha de generación", _fecha_local(inf["fecha_generacion"])),
            ("Consulta", inf["audit_trail"]["consulta_id"]), ("Aviso", LEYENDA)]
    for i, (k, v) in enumerate(meta, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=v)
        if k == "Techo aplicado":
            c.number_format = moneda
    fila0 = len(meta) + 4
    cabecera(ws, fila0, ["Indicador", "Valor"], [34, 44])
    r = inf["resumen_ejecutivo"]
    indicadores = [("Clientes consultados", r["clientes_consultados"], None),
                   ("Con saldo disponible", r["clientes_con_saldo"], None),
                   ("En el límite (≤ 5.000 €)", r["clientes_en_limite"], None),
                   ("Límite superado", r["clientes_superado_limite"], None),
                   ("Con datos pendientes", r["clientes_pendientes"], None),
                   ("Saldo disponible conjunto", r["total_saldo_disponible_conjunto"], moneda),
                   ("Total concedido en el periodo", r["total_concedido_periodo"], moneda),
                   ("Exceso sobre el techo", r["total_limite_superado"], moneda)]
    for i, (k, v, fmt) in enumerate(indicadores, start=fila0 + 1):
        ws.cell(row=i, column=1, value=k).border = borde
        c = ws.cell(row=i, column=2, value=v)
        c.border = borde
        c.alignment = Alignment(horizontal="right")
        if fmt:
            c.number_format = fmt
    ws.freeze_panes = None

    # Hoja 2 · Detalle
    wd = wb.create_sheet("Detalle")
    cabecera(wd, 1, ["CIF", "Razón social", "Saldo disponible", "Estado", "Años consultados", "Total concedido",
                     "Próxima salida del periodo", "Recupera saldo", "Pendientes", "Notas"],
             [12, 40, 16, 20, 18, 16, 16, 14, 50, 70])
    for i, d in enumerate(inf["detalle_clientes"], start=2):
        valores = [d["cif"], d["razon_social"], d["saldo_disponible"], d["estado_texto"],
                   ", ".join(map(str, d["anos_consultados"])), d["total_concedido"],
                   d["proxima_actualizacion_limite"], d["fecha_recupera_saldo"],
                   " ".join(d["pendientes"]), d["notas"]]
        for j, v in enumerate(valores, start=1):
            c = wd.cell(row=i, column=j, value=v)
            c.border = borde
            c.alignment = Alignment(vertical="top", wrap_text=j in (2, 9, 10))
            if j in (3, 6):
                c.number_format = moneda
    n = len(inf["detalle_clientes"])
    if n:
        tot = n + 2
        wd.cell(row=tot, column=2, value="TOTAL").font = Font(bold=True)
        for col in (3, 6):
            letra = get_column_letter(col)
            # el saldo negativo (límite superado) no resta del saldo conjunto, igual que en el resumen
            formula = f'=SUMIF({letra}2:{letra}{n + 1},">0")' if col == 3 else f"=SUM({letra}2:{letra}{n + 1})"
            c = wd.cell(row=tot, column=col, value=formula)
            c.number_format, c.font = moneda, Font(bold=True)
        rango = f"A2:J{n + 1}"
        for texto, color in (("Saldo disponible", "DCFCE7"), ("En el límite", "FEF3C7"),
                             ("Límite superado", "FEE2E2"), ("Pendiente de revisión", "E5E7EB")):
            wd.conditional_formatting.add(rango, FormulaRule(formula=[f'$D2="{texto}"'],
                                                             fill=PatternFill("solid", fgColor=color)))
    wd.auto_filter.ref = f"A1:J{max(n, 1) + 1}"

    # Hoja 3 · Ayudas computadas (como en Word y PDF: sin referencias de fuente ni ayudas que no computan)
    wr = wb.create_sheet("Ayudas computadas")
    cabecera(wr, 1, ["CIF", "Razón social", "Programa / convocatoria", "Nº convocatoria BDNS", "Código concesión",
                     "Concedente", "Reglamento", "Instrumento", "Año", "Fecha concesión", "Importe"],
             [12, 32, 50, 14, 16, 40, 30, 18, 7, 13, 15])
    fila = 2
    for d in inf["detalle_clientes"]:
        for x in _computan(d):
            valores = [d["cif"], d["razon_social"], x["programa"], x["numero_convocatoria"], x["codigo_concesion"],
                       x["convocante"], x["reglamento"], x["instrumento"], x["ano"], x["fecha_concesion"], x["importe"]]
            for j, v in enumerate(valores, start=1):
                c = wr.cell(row=fila, column=j, value=v)
                c.border = borde
                if j == 11:
                    c.number_format = moneda
            fila += 1
    wr.auto_filter.ref = f"A1:K{max(fila - 1, 1)}"

    wb.properties.creator = "MOVA IA · MOVA Minimis"
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def exportar(informe: dict, formato: str, carpeta: Path, base: str | None = None) -> Path:
    base = base or nombre_base(informe["informe"])
    funciones = {"docx": exportar_docx, "pdf": exportar_pdf, "xlsx": exportar_xlsx}
    return funciones[formato](informe, carpeta / f"{base}.{formato}")
