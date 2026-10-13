"""Agente 1 · Parser de entrada.

Saca pares (razón social, CIF) de documentos o de la entrada manual. Valida el
dígito de control de CIF, NIF y NIE: un identificador que no cuadra va a
`errores`, nunca se corrige a ojo. La razón social es secundaria: la oficial
la pone después la BDNS a partir del CIF.

Orden de lectura: texto nativo (PDF, Word, Excel, CSV, TXT) -> OCR con
Tesseract si está instalado -> Claude (visión) si hay clave configurada.
"""
import csv
import io
import re
import shutil
from pathlib import Path

LETRAS_NIF = "TRWAGMYFPDXBNJZSQVHLCKE"
LETRAS_CONTROL_CIF = "JABCDEFGHI"
_SEP = r"[\s.\-]?"
# Admite "B12345678", "B-12345678", "ES B12.345.678", "12345678Z", "X1234567L"
_PATRON_ID = re.compile(
    r"(?<![A-Z0-9])(?:ES[\s\-]?)?("
    r"[ABCDEFGHJNPQRSUVW]" + _SEP + r"\d{2}" + _SEP + r"\d{3}" + _SEP + r"\d{2}" + _SEP + r"[0-9A-J]"
    r"|\d{8}" + _SEP + r"[A-Z]"
    r"|[XYZ]" + _SEP + r"\d{7}" + _SEP + r"[A-Z]"
    r")(?![A-Z0-9])"
)
_ETIQUETAS = re.compile(r"\b(C\.?I\.?F\.?|N\.?I\.?F\.?|N\.?I\.?E\.?|CIF/NIF|RAZ[OÓ]N SOCIAL|EMPRESA|CLIENTE|"
                        r"BENEFICIARIO|NOMBRE)\b\s*[:\-]?", re.I)
_FORMAS = re.compile(r"\b(S\.?\s?L\.?\s?U?\.?|S\.?\s?A\.?\s?U?\.?|S\.?\s?COOP\.?|S\.?\s?L\.?\s?L\.?|C\.?B\.?|"
                     r"S\.?\s?C\.?|FUNDACI[OÓ]N|ASOCIACI[OÓ]N)\b", re.I)

EXTENSIONES = {
    ".pdf": "pdf", ".png": "imagen", ".jpg": "imagen", ".jpeg": "imagen", ".tif": "imagen", ".tiff": "imagen",
    ".bmp": "imagen", ".webp": "imagen", ".xlsx": "excel", ".xlsm": "excel", ".csv": "csv", ".txt": "txt",
    ".docx": "word",
}


# ---------------------------------------------------------------- identificadores
def validar_identificador(texto: str) -> tuple[str, str, bool, str]:
    """-> (normalizado, tipo CIF/NIF/NIE, válido, motivo)."""
    s = re.sub(r"[\s.\-]", "", (texto or "").upper())
    if s.startswith("ES") and len(s) == 11:
        s = s[2:]
    if re.fullmatch(r"\d{6,7}[A-Z]", s):  # DNI copiado sin ceros a la izquierda
        s = s.zfill(9)
    if re.fullmatch(r"[KLM]\d{7}[A-Z]", s):  # NIF especiales: letra como el DNI
        ok = LETRAS_NIF[int(s[1:8]) % 23] == s[8]
        return s, "NIF", ok, "" if ok else "letra de control del NIF incorrecta"
    if re.fullmatch(r"\d{8}[A-Z]", s):
        ok = LETRAS_NIF[int(s[:8]) % 23] == s[8]
        return s, "NIF", ok, "" if ok else "letra de control del NIF incorrecta"
    if re.fullmatch(r"[XYZ]\d{7}[A-Z]", s):
        num = int(str("XYZ".index(s[0])) + s[1:8])
        ok = LETRAS_NIF[num % 23] == s[8]
        return s, "NIE", ok, "" if ok else "letra de control del NIE incorrecta"
    if re.fullmatch(r"[ABCDEFGHJNPQRSUVW]\d{7}[0-9A-J]", s):
        d = s[1:8]
        pares = sum(int(d[i]) for i in (1, 3, 5))
        impares = sum(sum(divmod(int(d[i]) * 2, 10)) for i in (0, 2, 4, 6))
        digito = (10 - (pares + impares) % 10) % 10
        letra = LETRAS_CONTROL_CIF[digito]
        control = s[8]
        if s[0] in "PQRSNW":
            ok = control == letra
        elif s[0] in "ABEH":
            ok = control == str(digito)
        else:
            ok = control in (str(digito), letra)
        return s, "CIF", ok, "" if ok else "dígito de control del CIF incorrecto"
    return s, "", False, "formato no reconocido como CIF, NIF o NIE español"


def _limpiar_nombre(texto: str) -> str:
    t = _ETIQUETAS.sub(" ", texto or "")
    t = re.sub(r"[;|\t]+", " ", t)
    t = re.sub(r"\s{2,}", " ", t).strip(" ,:-–·")
    return t[:160]


def _cliente(cif, nombre, origen, confianza, notas="", anos=None):
    return {"razon_social": nombre or None, "cif": cif, "anos_solicitados": anos or [],
            "origen": origen, "confianza": round(confianza, 2), "notas": notas}


def _sumar(destino: dict, errores: list, cif_txt, nombre, origen, confianza, notas=""):
    cif, tipo, ok, motivo = validar_identificador(cif_txt)
    if not ok:
        errores.append({"entrada": f"{cif_txt} {nombre or ''}".strip(), "error": motivo})
        return
    nombre = _limpiar_nombre(nombre) if nombre else ""
    if tipo != "CIF":
        notas = (notas + " " if notas else "") + f"Identificador de tipo {tipo} (persona física)."
    if confianza < 0.8 and "Requiere revisión" not in notas:
        notas = (notas + " " if notas else "") + "Requiere revisión."
    previo = destino.get(cif)
    if previo:
        if not previo["razon_social"] and nombre:
            previo["razon_social"] = nombre
        previo["confianza"] = max(previo["confianza"], confianza)
        return
    destino[cif] = _cliente(cif, nombre, origen, confianza, notas.strip())


# ---------------------------------------------------------------- texto libre
def extraer_de_texto(texto: str, origen: str, destino: dict, errores: list, confianza_base=0.9) -> int:
    n = 0
    for linea in (texto or "").splitlines():
        linea_up = linea.upper()
        for m in _PATRON_ID.finditer(linea_up):
            resto = (linea[:m.start()] + " " + linea[m.end():])
            nombre = _limpiar_nombre(resto)
            conf = confianza_base if (nombre and _FORMAS.search(nombre)) else confianza_base - 0.25
            _sumar(destino, errores, m.group(1), nombre if len(nombre) > 2 else "", origen, conf)
            n += 1
        if not _PATRON_ID.search(linea_up):
            # Lo que parece un CIF pero tiene caracteres que el OCR confunde (7→/, 0→O, 1→l): se avisa, no se corrige
            for m in _PATRON_DUDOSO.finditer(linea_up):
                if sum(ch.isdigit() for ch in m.group(0)) < 5:
                    continue
                errores.append({"entrada": m.group(0).strip(),
                                "error": "posible identificador ilegible (revise el documento y escríbalo a mano)"})
    return n


_PATRON_DUDOSO = re.compile(r"(?<![A-Z0-9])[ABCDEFGHJNPQRSUVWXYZ][\s.\-]?[0-9OIL|/S]{7}[0-9A-Z](?![A-Z0-9])")


# ---------------------------------------------------------------- tablas (Excel, CSV)
_COL_CIF = re.compile(r"\b(cif|nif|nie|cif/nif|n\.?i\.?f|identificador|documento)\b", re.I)
_COL_NOMBRE = re.compile(r"(raz[oó]n social|nombre|empresa|cliente|beneficiario|denominaci[oó]n)", re.I)


def extraer_de_tabla(filas: list[list], origen: str, destino: dict, errores: list) -> int:
    filas = [[("" if c is None else str(c)).strip() for c in f] for f in filas if f and any(f)]
    if not filas:
        return 0
    cab = None
    for i, fila in enumerate(filas[:10]):
        ic = next((j for j, c in enumerate(fila) if _COL_CIF.search(c)), None)
        if ic is not None:
            inom = next((j for j, c in enumerate(fila) if j != ic and _COL_NOMBRE.search(c)), None)
            cab = (i, ic, inom)
            break
    if cab is None:  # sin cabecera reconocible: se trata como texto
        return extraer_de_texto("\n".join(" ; ".join(f) for f in filas), origen, destino, errores, 0.85)
    i0, ic, inom = cab
    n = 0
    for fila in filas[i0 + 1:]:
        if ic >= len(fila) or not fila[ic]:
            continue
        nombre = fila[inom] if inom is not None and inom < len(fila) else ""
        _sumar(destino, errores, fila[ic], nombre, origen, 0.97 if nombre else 0.85)
        n += 1
    return n


# ---------------------------------------------------------------- lectura de ficheros
def tesseract_disponible() -> str | None:
    import pytesseract
    candidatos = [shutil.which("tesseract"), r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                  r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"]
    for c in candidatos:
        if c and Path(c).exists():
            pytesseract.pytesseract.tesseract_cmd = c
            return c
    return None


def _ocr_imagenes(imagenes) -> str:
    import pytesseract
    textos = []
    for im in imagenes:
        try:
            textos.append(pytesseract.image_to_string(im, lang="spa"))
        except pytesseract.TesseractError:
            textos.append(pytesseract.image_to_string(im))  # sin paquete de idioma español
    return "\n".join(textos)


def _paginas_como_imagen(ruta: Path, max_paginas=20):
    import pymupdf as fitz
    from PIL import Image
    with fitz.open(ruta) as doc:
        for i, pag in enumerate(doc):
            if i >= max_paginas:
                break
            pix = pag.get_pixmap(dpi=200)
            yield Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def leer_texto(ruta: Path, tipo: str) -> tuple[str | list, str]:
    """-> (texto o filas, método). Las tablas vuelven como lista de filas."""
    if tipo == "excel":
        from openpyxl import load_workbook
        wb = load_workbook(ruta, read_only=True, data_only=True)
        filas = []
        for hoja in wb.worksheets:
            filas.extend(list(r) for r in hoja.iter_rows(values_only=True))
        wb.close()
        return filas, "excel"
    if tipo in ("csv", "txt"):
        crudo = ruta.read_bytes()
        for cod in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                texto = crudo.decode(cod)
                break
            except UnicodeDecodeError:
                continue
        if tipo == "csv":
            try:
                dialecto = csv.Sniffer().sniff(texto[:4000], delimiters=";,\t|")
            except csv.Error:
                dialecto = csv.excel
            return list(csv.reader(io.StringIO(texto), dialecto)), "csv"
        return texto, "texto"
    if tipo == "word":
        from docx import Document
        d = Document(ruta)
        partes = [p.text for p in d.paragraphs]
        for t in d.tables:
            for f in t.rows:
                partes.append(" ; ".join(c.text for c in f.cells))
        return "\n".join(partes), "word"
    if tipo == "pdf":
        from pypdf import PdfReader
        lector = PdfReader(ruta)
        texto = "\n".join((p.extract_text() or "") for p in lector.pages)
        if len(texto.strip()) >= 10 * max(1, len(lector.pages)):  # menos: PDF escaneado
            return texto, "pdf-texto"
        if tesseract_disponible():
            return _ocr_imagenes(_paginas_como_imagen(ruta)), "pdf-ocr"
        return texto, "pdf-sin-texto"
    if tipo == "imagen":
        if tesseract_disponible():
            from PIL import Image
            with Image.open(ruta) as im:
                return _ocr_imagenes([im.convert("RGB")]), "imagen-ocr"
        return "", "imagen-sin-ocr"
    raise ValueError(f"Tipo de fichero no admitido: {ruta.suffix}")


# ---------------------------------------------------------------- Claude como lector de respaldo
SISTEMA_CLAUDE = (
    "Eres el agente de extracción de MOVA Minimis. Lees documentos (a veces escaneados con mala calidad) "
    "y devuelves las empresas o personas que aparecen con su CIF, NIF o NIE español. Reglas: copia el "
    "identificador exactamente como aparece, sin corregir dígitos; si un carácter es ilegible, ponlo como '?' "
    "y explícalo en notas; no inventes razones sociales ni identificadores; si un dato no consta, déjalo vacío. "
    "Indica tu confianza de 0 a 1 por cada fila."
)
ESQUEMA_CLAUDE = {
    "type": "object",
    "properties": {
        "clientes": {"type": "array", "items": {
            "type": "object",
            "properties": {
                "razon_social": {"type": "string"},
                "cif": {"type": "string"},
                "confianza": {"type": "number"},
                "notas": {"type": "string"},
            },
            "required": ["razon_social", "cif", "confianza", "notas"],
            "additionalProperties": False,
        }},
    },
    "required": ["clientes"],
    "additionalProperties": False,
}


def _bloques_para_claude(ruta: Path, tipo: str, texto: str) -> list:
    import base64
    if tipo == "pdf":
        datos = base64.standard_b64encode(ruta.read_bytes()).decode()
        return [{"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": datos}},
                {"type": "text", "text": "Extrae los identificadores fiscales y razones sociales de este documento."}]
    if tipo == "imagen":
        from PIL import Image
        with Image.open(ruta) as im:
            im = im.convert("RGB")
            im.thumbnail((2000, 2000))
            buf = io.BytesIO()
            im.save(buf, "PNG")
        datos = base64.standard_b64encode(buf.getvalue()).decode()
        return [{"type": "image", "source": {"type": "base64", "media_type": "image/png", "data": datos}},
                {"type": "text", "text": "Extrae los identificadores fiscales y razones sociales de esta imagen."}]
    return [{"type": "text", "text": "Extrae los identificadores fiscales y razones sociales de este texto:\n\n"
                                      + texto[:150_000]}]


# ---------------------------------------------------------------- punto de entrada del agente
def procesar_documento(ruta: Path, nombre_original: str, llamar_claude=None, aviso=None) -> dict:
    """`llamar_claude(sistema, bloques, esquema)` lo inyecta la capa de rutas si hay clave configurada."""
    tipo = EXTENSIONES.get(Path(nombre_original).suffix.lower())
    destino, errores, avisos = {}, [], []
    if not tipo:
        return {"clientes": [], "errores": [{"entrada": nombre_original, "error": "tipo de fichero no admitido"}],
                "avisos": [], "metodo": None, "tipo": None}
    contenido, metodo = leer_texto(ruta, tipo)
    origen = f"documento:{nombre_original}"
    if isinstance(contenido, list):
        extraer_de_tabla(contenido, origen, destino, errores)
        texto_plano = "\n".join(" ; ".join("" if c is None else str(c) for c in f) for f in contenido)
    else:
        extraer_de_texto(contenido, origen, destino, errores)
        texto_plano = contenido

    sin_texto = metodo in ("pdf-sin-texto", "imagen-sin-ocr")
    if not destino and llamar_claude and (sin_texto or texto_plano.strip()):
        if aviso:
            aviso(f"{nombre_original}: sin identificadores legibles por lectura directa; se pide a Claude.")
        respuesta = llamar_claude(SISTEMA_CLAUDE, _bloques_para_claude(ruta, tipo, texto_plano), ESQUEMA_CLAUDE)
        for c in respuesta.get("clientes", []):
            conf = float(c.get("confianza") or 0.5)
            _sumar(destino, errores, c.get("cif", ""), c.get("razon_social", ""), origen + " (Claude)",
                   min(conf, 0.9), c.get("notas", ""))
        metodo += "+claude"
    elif not destino and sin_texto:
        avisos.append(f"{nombre_original}: OCR no disponible. Introduzca los datos manualmente o copie y pegue "
                      "el texto del documento (o instale Tesseract, o configure la clave de Claude).")
    elif not destino:
        avisos.append(f"{nombre_original}: no se encontró ningún CIF, NIF o NIE válido.")
    return {"clientes": list(destino.values()), "errores": errores, "avisos": avisos, "metodo": metodo,
            "tipo": tipo}


def procesar_manual(texto: str) -> dict:
    """Entrada manual: una línea por cliente, 'CIF ; razón social' en cualquier orden."""
    destino, errores = {}, []
    for linea in (texto or "").splitlines():
        if not linea.strip():
            continue
        if not _PATRON_ID.search(linea.upper()):
            errores.append({"entrada": linea.strip()[:120],
                            "error": "no contiene un CIF, NIF o NIE con formato español"})
            continue
        extraer_de_texto(linea, "manual", destino, errores, confianza_base=1.0)
    for c in destino.values():  # lo tecleado por el usuario no necesita revisión por confianza
        c["confianza"] = 1.0
        c["notas"] = c["notas"].replace("Requiere revisión.", "").strip()
    return {"clientes": list(destino.values()), "errores": errores, "avisos": [], "metodo": "manual"}


def unir(*salidas: dict, anos: list[int] | None = None) -> dict:
    clientes, errores, avisos = {}, [], []
    for s in salidas:
        for c in s.get("clientes", []):
            previo = clientes.get(c["cif"])
            if previo is None:
                clientes[c["cif"]] = dict(c)
            elif not previo.get("razon_social") and c.get("razon_social"):
                previo["razon_social"] = c["razon_social"]
        errores += s.get("errores", [])
        avisos += s.get("avisos", [])
    for c in clientes.values():
        c["anos_solicitados"] = anos or []
    return {"clientes": list(clientes.values()), "errores": errores, "avisos": avisos}
