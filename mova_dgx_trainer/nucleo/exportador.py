"""Exportador: informe Word (python-docx) y PDF (reportlab) con el mismo contenido.

Secciones: portada, resumen con KPIs, dataset y procedencia, configuración e
hiperparámetros, ejecuciones y reintentos con veredicto, gráficas y anexo de
trazabilidad. Lo que falte se escribe como [PENDIENTE: dato].
"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd

from nucleo import graficas
from nucleo.contexto import pendiente

TITULO = "MOVA DGX Spark Trainer"


def _v(valor, dato: str, formato: str | None = None) -> str:
    if valor is None or valor == "" or valor == {} or valor == []:
        return pendiente(dato)
    if formato and isinstance(valor, (int, float)):
        return format(valor, formato)
    if isinstance(valor, (dict, list)):
        return json.dumps(valor, ensure_ascii=False)
    return str(valor)


def generar_graficas(datos: dict, carpeta: Path) -> list[tuple[str, Path]]:
    """Devuelve [(título, ruta_png)]."""
    salida: list[tuple[str, Path]] = []
    nombre = datos["kpis"].get("metrica_nombre") or "métrica"
    filas = [{"etiqueta": f"#{e['id']}", "valor": e.get("metrica")}
             for e in datos.get("ejecuciones", []) if e.get("estado") == "ok"]
    if filas:
        fig = graficas.comparacion_ejecuciones(filas, datos.get("linea_base", {}).get(nombre), nombre)
        salida.append(("Ejecuciones frente a la línea base",
                       graficas.guardar(fig, carpeta / "comparacion.png")))
    mejor = datos.get("mejor") or {}
    if mejor.get("curva"):
        salida.append(("Curva de entrenamiento de la mejor ejecución",
                       graficas.guardar(graficas.curva_entrenamiento(mejor["curva"]),
                                        carpeta / "curva.png")))
    metricas = datos.get("metricas_mejor") or {}
    if metricas.get("matriz_confusion"):
        salida.append(("Matriz de confusión",
                       graficas.guardar(graficas.matriz_confusion(metricas["matriz_confusion"]),
                                        carpeta / "matriz_confusion.png")))
    ruta_pred = datos.get("ruta_predicciones")
    if ruta_pred and Path(ruta_pred).exists() and not metricas.get("matriz_confusion"):
        p = pd.read_csv(ruta_pred, encoding="utf-8")
        fig = graficas.real_frente_predicho(p["y_real"].astype(float).tolist(),
                                            p["y_pred"].astype(float).tolist())
        salida.append(("Real frente a predicho", graficas.guardar(fig, carpeta / "real_predicho.png")))
    return salida


def _secciones(datos: dict) -> list[tuple[str, list]]:
    """Contenido común a Word y PDF: [(título, [bloques])].

    Bloque: ("p", texto) | ("tabla", cabecera, filas) | ("img", título, ruta)
    """
    k, p, ds = datos["kpis"], datos["proyecto"], datos.get("dataset") or {}
    cfg, traza = datos.get("configuracion") or {}, datos.get("trazabilidad") or {}
    base = datos.get("linea_base") or {}
    nombre = k.get("metrica_nombre") or "métrica"
    secciones = []

    resumen = [("p", f"Proyecto «{p['nombre']}» · tipo de tarea {p['tipo_tarea']}."),
               ("tabla", ["Indicador", "Valor"], [
                   [nombre, _v(k.get("metrica_valor"), nombre, ".4f")],
                   ["Mejora sobre baseline", _v(k.get("mejora"), "mejora", "+.4f")],
                   ["Duración (s)", _v(k.get("duracion_s"), "duración", ".1f")],
                   ["Reintentos", _v(k.get("reintentos"), "reintentos")],
                   ["Veredicto", _v(k.get("veredicto"), "veredicto")],
               ])]
    if datos.get("mejor"):
        resumen.append(("p", f"Mejor ejecución: #{datos['mejor']['id']} "
                             f"(intento {datos['mejor']['intento']})."))
    secciones.append(("Resumen", resumen))

    filas_base = [[m, _v(v, m, ".4f")] for m, v in base.items()
                  if isinstance(v, (int, float)) and not isinstance(v, bool)]
    secciones.append(("Línea base", [
        ("p", f"Estrategia: {_v(base.get('estrategia'), 'estrategia de la línea base')}, "
              "sobre la misma división que el modelo."),
        ("tabla", ["Métrica", "Valor"], filas_base or [[pendiente("línea base"), ""]])]))

    division = ds.get("division") or {}
    ficheros = datos.get("ficheros") or []
    secciones.append(("Dataset y procedencia", [
        ("tabla", ["Campo", "Valor"], [
            ["Filas / columnas", f"{_v(ds.get('filas'), 'filas')} / {_v(ds.get('columnas'), 'columnas')}"],
            ["Objetivo", _v(ds.get("objetivo"), "objetivo")],
            ["División", _v(division.get("filas"), "división")],
            ["Estratificada", _v(division.get("estratificada"), "estratificación")],
            ["Hash del dataset (SHA-256)", _v(ds.get("hash"), "hash del dataset")],
        ]),
        ("tabla", ["Fichero", "Formato", "Filas", "Estado", "Hash (SHA-256)"],
         [[f.get("nombre"), _v(f.get("formato"), "formato"), _v(f.get("filas"), "filas"),
           f.get("estado"), _v((f.get("hash") or "")[:16], "hash")] for f in ficheros]
         or [[pendiente("ficheros de origen"), "", "", "", ""]]),
    ]))

    hp = cfg.get("hiperparametros") or {}
    secciones.append(("Configuración e hiperparámetros", [
        ("tabla", ["Parámetro", "Valor"], [
            ["Framework", _v(cfg.get("framework"), "framework")],
            ["Script", _v(cfg.get("script"), "script")],
            ["Métrica objetivo", _v(cfg.get("metrica_objetivo"), "métrica objetivo")],
            ["Umbrales", _v(cfg.get("umbrales"), "umbrales")],
            ["Máximo de reintentos", _v(cfg.get("max_reintentos"), "máximo de reintentos")],
            ["Semilla", _v(cfg.get("semilla"), "semilla")],
        ] + [[f"hp · {c}", _v(v, c)] for c, v in hp.items()]),
    ]))

    filas_ej = [[f"#{e['id']}", str(e.get("intento")), e.get("motivo") or "inicial",
                 e.get("estado") or "", e.get("veredicto") or pendiente("veredicto"),
                 _v(e.get("metrica"), nombre, ".4f"), _v(e.get("duracion_s"), "duración", ".1f")]
                for e in datos.get("ejecuciones", [])]
    secciones.append(("Ejecuciones y reintentos", [
        ("tabla", ["Ejecución", "Intento", "Motivo", "Estado", "Veredicto", nombre, "Duración (s)"],
         filas_ej or [[pendiente("ejecuciones"), "", "", "", "", "", ""]])]))

    imgs = [("img", t, r) for t, r in datos.get("graficas", [])]
    secciones.append(("Gráficas", imgs or [("p", pendiente("gráficas"))]))

    versiones = traza.get("versiones") or {}
    secciones.append(("Anexo de trazabilidad", [
        ("tabla", ["Elemento", "Valor"], [
            ["Semilla", _v(traza.get("semilla"), "semilla")],
            ["Hash del dataset", _v(traza.get("hash_dataset"), "hash del dataset")],
            ["Modo de ejecución", _v(traza.get("modo"), "modo")],
            ["Comando", _v(traza.get("comando"), "comando")],
            ["Log íntegro", _v(traza.get("ruta_log"), "ruta del log")],
            ["Modelo", _v(traza.get("ruta_modelo"), "ruta del modelo")],
        ] + [[f"versión · {c}", _v(v, c)] for c, v in versiones.items()]
          + [[f"hash · {n}", _v(h, "hash")] for n, h in (traza.get("hashes_ficheros") or {}).items()]),
    ]))
    return secciones


def exportar_docx(datos: dict, ruta: Path) -> Path:
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(10)
    doc.add_heading(TITULO, level=0)
    doc.add_paragraph(f"Informe de entrenamiento · {datos['proyecto']['nombre']}")
    doc.add_paragraph(f"Tipo de tarea: {datos['proyecto']['tipo_tarea']}")
    doc.add_paragraph(f"Generado: {datos['generado']}")
    doc.add_page_break()
    for titulo, bloques in _secciones(datos):
        doc.add_heading(titulo, level=1)
        for b in bloques:
            if b[0] == "p":
                par = doc.add_paragraph(b[1])
                if "[PENDIENTE" in b[1]:
                    for r in par.runs:
                        r.font.color.rgb = RGBColor(0xD9, 0x77, 0x06)
            elif b[0] == "tabla":
                tabla = doc.add_table(rows=1, cols=len(b[1]))
                tabla.style = "Light Grid Accent 1"
                for i, c in enumerate(b[1]):
                    tabla.rows[0].cells[i].text = str(c)
                for fila in b[2]:
                    celdas = tabla.add_row().cells
                    for i, v in enumerate(fila):
                        celdas[i].text = str(v)
                        if "[PENDIENTE" in str(v):
                            for r in celdas[i].paragraphs[0].runs:
                                r.font.color.rgb = RGBColor(0xD9, 0x77, 0x06)
            elif b[0] == "img":
                doc.add_paragraph(b[1])
                doc.add_picture(str(b[2]), width=Cm(15))
    ruta.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(ruta))
    return ruta


def _escapar(texto: str) -> str:
    return (str(texto).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def exportar_pdf(datos: dict, ruta: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)

    estilos = getSampleStyleSheet()
    celda = estilos["BodyText"].clone("celda", fontSize=8, leading=10)
    ambar = "#d97706"

    def parrafo(texto, estilo=celda):
        t = _escapar(texto)
        if "[PENDIENTE" in t:
            t = f'<font color="{ambar}">{t}</font>'
        return Paragraph(t, estilo)

    historia = [Spacer(1, 5 * cm), Paragraph(TITULO, estilos["Title"]),
                Paragraph(f"Informe de entrenamiento · {_escapar(datos['proyecto']['nombre'])}",
                          estilos["Heading2"]),
                Paragraph(f"Tipo de tarea: {_escapar(datos['proyecto']['tipo_tarea'])}",
                          estilos["BodyText"]),
                Paragraph(f"Generado: {_escapar(datos['generado'])}", estilos["BodyText"]),
                PageBreak()]
    ancho = A4[0] - 4 * cm
    for titulo, bloques in _secciones(datos):
        historia.append(Paragraph(_escapar(titulo), estilos["Heading1"]))
        for b in bloques:
            if b[0] == "p":
                historia.append(parrafo(b[1], estilos["BodyText"]))
            elif b[0] == "tabla":
                filas = [[parrafo(c) for c in b[1]]] + [[parrafo(v) for v in f] for f in b[2]]
                t = Table(filas, colWidths=[ancho / len(b[1])] * len(b[1]), repeatRows=1)
                t.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0d98c")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa0a6")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP")]))
                historia.append(t)
            elif b[0] == "img":
                historia.append(Paragraph(_escapar(b[1]), estilos["Heading3"]))
                historia.append(Image(str(b[2]), width=15 * cm, height=15 * cm * 0.5625,
                                      kind="proportional"))
            historia.append(Spacer(1, 0.3 * cm))
    ruta.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(ruta), pagesize=A4, title=TITULO, author="MOVA",
                      leftMargin=2 * cm, rightMargin=2 * cm).build(historia)
    return ruta


def exportar_informe(datos: dict, carpeta: Path, nombre_base: str) -> tuple[Path, Path]:
    carpeta.mkdir(parents=True, exist_ok=True)
    datos = dict(datos)
    datos["graficas"] = generar_graficas(datos, carpeta / "graficas")
    return (exportar_docx(datos, carpeta / f"{nombre_base}.docx"),
            exportar_pdf(datos, carpeta / f"{nombre_base}.pdf"))


def exportar_tabla(cabecera: list[str], filas: list[list], ruta: Path) -> Path:
    """Exporta una tabla visible a Excel (.xlsx) o CSV según la extensión."""
    df = pd.DataFrame(filas, columns=cabecera)
    ruta = Path(ruta)
    if ruta.suffix.lower() == ".xlsx":
        df.to_excel(ruta, index=False)
    else:
        df.to_csv(ruta, index=False, encoding="utf-8-sig", sep=";")
    return ruta
